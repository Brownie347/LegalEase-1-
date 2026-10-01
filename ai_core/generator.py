"""Formatting utilities: sanitise AI text and export it as HTML / DOCX / PDF.

Public API
    sanitize_text(text)                      -> str
    format_html_preview(text)                -> str
    format_docx(text, doc_type, ...)         -> bytes
    format_pdf(text, doc_type, ...)          -> bytes
"""
from __future__ import annotations

import html
import io
import re
from pathlib import Path
from typing import List, Optional, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF
from PIL import Image

import config

# =========================================================================== #
# Text clean-up
# =========================================================================== #
_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201a": "'", "\u201b": "'",
    "\u201c": '"', "\u201d": '"', "\u201e": '"',
    "\u2013": "-", "\u2014": "-", "\u2212": "-", "\u2010": "-", "\u2011": "-",
    "\u2026": "...", "\u00a0": " ", "\u200b": "", "\ufeff": "",
    "\u2022": "-", "\u25cf": "-", "\u25aa": "-", "\u2023": "-",
    "\u20b9": "Rs. ", "\u20ac": "EUR ",
}
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def sanitize_text(text: str) -> str:
    """Remove markdown symbols, typographic quotes and control characters."""
    if not text:
        return ""
    for src, dst in _REPLACEMENTS.items():
        text = text.replace(src, dst)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL.sub("", text)
    text = re.sub(r"```[a-zA-Z]*", "", text).replace("`", "")
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)                       # **bold**
    text = re.sub(r"(?<![\w*])\*([^*\n]+)\*(?![\w*])", r"\1", text)     # *italic*
    text = re.sub(r"^[ \t]{0,3}#{1,6}[ \t]*", "", text, flags=re.M)     # ## headings
    text = re.sub(r"^[ \t]*[*+][ \t]+", "- ", text, flags=re.M)         # * bullets
    text = re.sub(r"_{31,}", "_" * 30, text)                            # signature lines
    text = re.sub(r"[ \t]+$", "", text, flags=re.M)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _latin1(text: str) -> str:
    """FPDF core fonts only support latin-1."""
    return text.encode("latin-1", "replace").decode("latin-1")


# =========================================================================== #
# Line classification (shared by HTML / DOCX / PDF)
# =========================================================================== #
_NUMBERED = re.compile(
    r"^(?:\d+(?:\.\d+)*[.)]|(?:Section|Article|Clause)\s+[\dIVXLC]+[.:)]?)\s+\S", re.I
)
_BULLET = re.compile(
    r"^(?:[-•]\s+|\(?[a-zA-Z]\)\s+|\((?:i{1,3}|iv|v|vi{0,3}|ix|x)\)\s+|\(\d{1,2}\)\s+)"
)


def classify_line(line: str) -> str:
    """Return 'blank', 'heading', 'bullet' or 'para'."""
    s = line.strip()
    if not s:
        return "blank"
    if _BULLET.match(s):
        return "bullet"
    letters = [c for c in s if c.isalpha()]
    words = len(s.split())
    if letters and all(c.isupper() for c in letters) and len(s) <= 90 and not s.endswith(","):
        return "heading"
    if s.endswith(":") and len(s) <= 70 and words <= 9:
        return "heading"
    if _NUMBERED.match(s) and len(s) <= 70 and not s.endswith((".", ";", ",")):
        return "heading"
    return "para"


def _split_title(lines: List[str]) -> Tuple[Optional[str], List[str]]:
    """Pull a title-looking first line off the body, if there is one."""
    for i, raw in enumerate(lines):
        first = raw.strip()
        if not first:
            continue
        looks_like_title = (
            len(first) <= 90
            and not first.endswith((".", ";", ",", ":"))
            and classify_line(first) != "bullet"
            and not _NUMBERED.match(first)
        )
        if looks_like_title:
            return first, lines[i + 1:]
        return None, lines[i:]
    return None, []


def parse_terms(terms: str) -> List[str]:
    """Split the semicolon/newline separated Terms field into clean items."""
    items = []
    for part in re.split(r"[;\n]", terms or ""):
        part = re.sub(r"^\s*(?:[-•*]|\d+[.)])\s*", "", part).strip()
        if part:
            items.append(part)
    return items


# =========================================================================== #
# HTML preview
# =========================================================================== #
_PALETTES = {
    "dark": {"title": "#ffffff", "heading": "#8ab4f8", "text": "#dfe4ee"},
    "paper": {"title": "#14203A", "heading": "#14203A", "text": "#23262F"},
}


def format_html_preview(text: str, theme: str = "dark") -> str:
    """Convert plain text to styled HTML blocks (single line, no blank lines,
    so it is safe to pass to st.markdown(..., unsafe_allow_html=True)).

    theme: "dark" (light text for a dark card) or "paper" (dark text for a white sheet).
    """
    c = _PALETTES.get(theme, _PALETTES["dark"])
    title, body = _split_title(sanitize_text(text).split("\n"))
    out: List[str] = []

    def esc(s: str) -> str:
        return html.escape(s).replace("$", "&#36;")  # '$' would trigger LaTeX in Streamlit

    if title:
        out.append(
            f"<div style='text-align:center;color:{c['title']};font-size:1.45em;font-weight:700;"
            f"margin:0 0 1.1em 0;line-height:1.25'>{esc(title)}</div>"
        )
    for line in body:
        kind = classify_line(line)
        s = esc(line.strip())
        if kind == "blank":
            continue
        if kind == "heading":
            out.append(
                f"<div style='color:{c['heading']};font-weight:700;margin:1.5em 0 .35em 0;"
                f"font-size:1.02em'>{s}</div>"
            )
        elif kind == "bullet":
            out.append(f"<p style='margin:.3em 0 .3em 1.8em;color:{c['text']}'>{s}</p>")
        else:
            out.append(f"<p style='margin:.55em 0;color:{c['text']};text-align:justify'>{s}</p>")
    return "".join(out)


def document_stats(text: str) -> dict:
    """Word count, section count and a rough page estimate for the UI."""
    clean = sanitize_text(text)
    lines = clean.split("\n")
    words = len(clean.split())
    numbered = sum(1 for l in lines if re.match(r"^\d+[.)]\s", l.strip()))
    headings = sum(1 for l in lines if classify_line(l) == "heading")
    return {
        "words": words,
        "sections": numbered or headings,
        "pages": max(1, round(words / 450)),
    }


# =========================================================================== #
# DOCX
# =========================================================================== #
_FONT = "Times New Roman"


def _style_run(run, size=12, bold=None, italic=None, color=None):
    run.font.name = _FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def _shade(cell, fill="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _add_page_number(paragraph):
    """Insert a PAGE field (each field part in its own styled run)."""
    def field_run():
        run = paragraph.add_run()
        _style_run(run, size=9, italic=True, color=(110, 110, 110))
        return run

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    field_run()._r.append(begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    field_run()._r.append(instr)

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    field_run()._r.append(separate)

    result = field_run()
    result.text = "1"

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run()._r.append(end)


def _set_col_widths(table, widths):
    """Fixed column widths that Word *and* LibreOffice respect."""
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def _table_row(table, left: str, right: str, bold_left=True, shade_left=True):
    cells = table.add_row().cells
    for cell, value, is_bold in ((cells[0], left, bold_left), (cells[1], right, False)):
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(2)
        _style_run(para.add_run(value), size=11, bold=is_bold)
    if shade_left:
        _shade(cells[0], "F2F2F2")
    return cells


def format_docx(
    text: str,
    doc_type: str,
    terms: str = "",
    parties: str = "",
    effective_date: str = "",
    logo_path: Optional[Path] = None,
) -> bytes:
    """Build a Word document: logo, title, summary + terms tables, body, footer."""
    text = sanitize_text(text)
    title, body = _split_title(text.split("\n"))
    title = title or (doc_type or "Legal Document").strip().title()
    logo = Path(logo_path) if logo_path else config.LOGO_PATH

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = _FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), _FONT)
    normal.font.size = Pt(12)

    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(1)

    # Logo -------------------------------------------------------------- #
    if logo.exists():
        doc.add_picture(str(logo), width=Inches(2.3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Title ------------------------------------------------------------- #
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(12)
    _style_run(p.add_run(title.upper()), size=16, bold=True, color=(27, 42, 73))

    # Summary + terms tables ------------------------------------------- #
    term_items = parse_terms(terms)
    if parties.strip() or effective_date.strip():
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_row(table, "Document Type", (doc_type or title).strip())
        if parties.strip():
            _table_row(table, "Parties", parties.strip())
        if effective_date.strip():
            _table_row(table, "Effective Date", effective_date.strip())
        _set_col_widths(table, (Inches(1.7), Inches(4.8)))
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    if term_items:
        h = doc.add_paragraph()
        _style_run(h.add_run("Key Terms"), size=12, bold=True)
        h.paragraph_format.space_after = Pt(4)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, label in zip(table.rows[0].cells, ("No.", "Term / Condition")):
            _style_run(cell.paragraphs[0].add_run(label), size=11, bold=True)
            _shade(cell)
        for i, item in enumerate(term_items, 1):
            _table_row(table, str(i), item, bold_left=False, shade_left=False)
        _set_col_widths(table, (Inches(0.6), Inches(5.9)))
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Body -------------------------------------------------------------- #
    for line in body:
        kind = classify_line(line)
        if kind == "blank":
            continue
        para = doc.add_paragraph()
        fmt = para.paragraph_format
        fmt.space_after = Pt(6)
        fmt.line_spacing = 1.15
        s = line.strip()
        if kind == "heading":
            fmt.space_before = Pt(12)
            fmt.keep_with_next = True
            _style_run(para.add_run(s), size=12, bold=True)
        elif kind == "bullet":
            fmt.left_indent = Inches(0.5)
            _style_run(para.add_run(s), size=12)
        else:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _style_run(para.add_run(s), size=12)

    # Footer (every page) ---------------------------------------------- #
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(footer_p.add_run(f"{config.FOOTER_TEXT}   |   Page "), size=9, italic=True, color=(110, 110, 110))
    _add_page_number(footer_p)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# =========================================================================== #
# PDF
# =========================================================================== #
class LegalPDF(FPDF):
    """A4 PDF with the logo + title on every page and a branded footer."""

    def __init__(self, doc_title: str, logo_path: Path, footer_text: str):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.doc_title = _latin1(doc_title)
        self.logo_path = Path(logo_path)
        self.footer_text = _latin1(footer_text)
        self.set_margins(20, 20, 20)
        self.set_auto_page_break(auto=True, margin=22)
        self.alias_nb_pages()
        self._logo_w = 44.0
        self._logo_h = 0.0
        if self.logo_path.exists():
            with Image.open(self.logo_path) as im:
                self._logo_h = self._logo_w * im.height / im.width

    def header(self):
        y = 10
        if self._logo_h:
            self.image(str(self.logo_path), x=(self.w - self._logo_w) / 2, y=y, w=self._logo_w)
            y += self._logo_h + 2
        self.set_y(y)
        self.set_font("Helvetica", "B", 11)
        self.set_text_color(27, 42, 73)
        self.cell(0, 7, self.doc_title, align="C", new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(200, 200, 200)
        self.line(self.l_margin, self.get_y() + 1, self.w - self.r_margin, self.get_y() + 1)
        self.set_y(self.get_y() + 6)
        self.set_text_color(0, 0, 0)

    def footer(self):
        self.set_y(-17)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(110, 110, 110)
        self.cell(0, 5, self.footer_text, align="C", new_x="LMARGIN", new_y="NEXT")
        self.cell(0, 5, f"Page {self.page_no()} of {{nb}}", align="C")
        self.set_text_color(0, 0, 0)


def format_pdf(
    text: str,
    doc_type: str,
    terms: str = "",
    parties: str = "",
    effective_date: str = "",
    logo_path: Optional[Path] = None,
) -> bytes:
    """Build a branded PDF: header logo, bold headings, bullet terms, footer."""
    text = sanitize_text(text)
    title, body = _split_title(text.split("\n"))
    title = title or (doc_type or "Legal Document").strip().title()
    logo = Path(logo_path) if logo_path else config.LOGO_PATH

    pdf = LegalPDF(title, logo, config.FOOTER_TEXT)
    pdf.add_page()

    def para(txt, style="", size=10.5, h=5.6, align="L", indent=0.0, gap=1.5):
        pdf.set_font("Helvetica", style, size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(0, h, _latin1(txt), align=align, new_x="LMARGIN", new_y="NEXT")
        if gap:
            pdf.ln(gap)

    # Summary block ------------------------------------------------------ #
    summary = []
    if parties.strip():
        summary.append(("Parties", parties.strip()))
    if effective_date.strip():
        summary.append(("Effective Date", effective_date.strip()))
    for label, value in summary:
        para(label, "B", 10.5, gap=0)
        para(value, "", 10.5)

    term_items = parse_terms(terms)
    if term_items:
        pdf.ln(1)
        para("Key Terms", "B", 11.5, gap=1)
        for item in term_items:
            para(f"-  {item}", "", 10.5, indent=4, gap=0.8)
    if summary or term_items:
        pdf.ln(1)
        pdf.set_draw_color(200, 200, 200)
        pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
        pdf.ln(4)

    # Body --------------------------------------------------------------- #
    for line in body:
        kind = classify_line(line)
        s = line.strip()
        if kind == "blank":
            continue
        if kind == "heading":
            pdf.ln(2.5)
            para(s, "B", 11.5, h=6, gap=1)
        elif kind == "bullet":
            para(s, "", 10.5, indent=8, gap=1)
        else:
            para(s, "", 10.5, align="J", gap=2)

    return bytes(pdf.output())
