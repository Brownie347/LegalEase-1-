import io

from docx import Document

from ai_core.generator import (
    classify_line,
    document_stats,
    format_docx,
    format_html_preview,
    format_pdf,
    parse_terms,
    sanitize_text,
)

SAMPLE = """## Freelance Work Contract

Agreement made this **15th** day of April, 2025. Fee: Rs. 5,000 and $100.

Between:

Jane Doe (the "Service Provider") and TechNova Inc. (the "Client").

1. Services:

The Service Provider shall deliver the work by May 15, 2025 and this line is long enough to be a paragraph.

2. Termination:

a) Mutual written agreement;
b) Breach by either party.

______________________________________________
Jane Doe (Service Provider)
"""
TERMS = "Work due May 15; Payment within 7 days; IP retained by client"


def test_sanitize_removes_markdown_and_fancy_quotes():
    out = sanitize_text("## Title\n\n**Bold** \u201cquoted\u201d \u2013 dash \u20b9500\n* item")
    assert "#" not in out and "**" not in out
    assert '"quoted"' in out and "-" in out and "Rs." in out
    assert "- item" in out


def test_classify_line():
    assert classify_line("") == "blank"
    assert classify_line("1. Services:") == "heading"
    assert classify_line("WITNESSETH:") == "heading"
    assert classify_line("a) Mutual written agreement;") == "bullet"
    assert classify_line("The Service Provider agrees to provide services to the Client.") == "para"


def test_parse_terms():
    assert parse_terms("A; B ;\n- C;") == ["A", "B", "C"]
    assert parse_terms("") == []


def test_html_preview_is_single_block_and_escapes():
    out = format_html_preview("Title\n\n1. Heading:\n\n<script>x</script> costs $5")
    assert "\n" not in out
    assert "<script>" not in out and "&lt;script&gt;" in out
    assert "$" not in out


def test_docx_contains_tables_footer_and_logo():
    data = format_docx(SAMPLE, "Freelance Work Contract", TERMS, "Jane Doe, TechNova Inc.", "April 15, 2025")
    doc = Document(io.BytesIO(data))
    assert len(doc.tables) == 2
    assert doc.tables[1].rows[1].cells[1].text == "Work due May 15"
    assert "All Rights Reserved" in doc.sections[0].footer.paragraphs[0].text
    assert len(doc.inline_shapes) == 1  # logo


def test_pdf_is_valid_and_multipage():
    data = format_pdf(SAMPLE * 12, "Freelance Work Contract", TERMS, "Jane Doe", "April 15, 2025")
    assert data.startswith(b"%PDF")
    assert data.count(b"/Type /Page\n") + data.count(b"/Type /Page ") >= 2 or len(data) > 5000


def test_pdf_handles_non_latin_text():
    data = format_pdf("Title\n\nCaf\u00e9 \u4e2d\u6587 \u0939\u093f\u0902\u0926\u0940", "Test")
    assert data.startswith(b"%PDF")


def test_html_preview_paper_theme_and_no_heading_tags():
    out = format_html_preview("Title\n\n1. Heading:\n\nBody text that is long enough to be a paragraph.", theme="paper")
    assert "#14203A" in out and "<h3" not in out and "<h4" not in out


def test_document_stats():
    stats = document_stats(SAMPLE)
    assert stats["sections"] == 2 and stats["pages"] == 1 and stats["words"] > 20
